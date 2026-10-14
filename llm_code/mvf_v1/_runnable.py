import subprocess, sys
subprocess.check_call([sys.executable, "-m", "pip", "install", "-q",
                       "yfinance", "openpyxl", "pillow", "requests", "scipy", "matplotlib", "pandas", "numpy"])

import io, json, zipfile, csv
import numpy as np
import pandas as pd
import requests
import yfinance as yf
import matplotlib
import matplotlib.pyplot as plt
from scipy.optimize import minimize
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment

# ----------------------------------------------------------------------------
# DATA
# ----------------------------------------------------------------------------
TICKERS = ["SPY", "IWM", "VEA", "VWO", "XLE", "QQQ", "AGG", "TIP", "HYG", "GLD"]
N = len(TICKERS)

raw = yf.download(TICKERS, start="2020-11-20", end="2026-01-10",
                  auto_adjust=False, progress=False)
adj = raw["Adj Close"][TICKERS].dropna(how="all")
adj.index = pd.to_datetime(adj.index).tz_localize(None)

# month-end using the LAST available trading day of each month
try:
    monthly = adj.resample("ME").last()
except ValueError:
    monthly = adj.resample("M").last()
monthly = monthly.dropna(how="all")
monthly = monthly.loc["2020-12-31":"2024-12-31"]
assert len(monthly) == 49, f"Expected 49 month-end prices, got {len(monthly)}"
assert not monthly.isna().any().any(), "Missing prices in monthly data"

rets = monthly / monthly.shift(1) - 1.0
rets = rets.iloc[1:]
assert len(rets) == 48, f"Expected 48 monthly returns, got {len(rets)}"

# Fama-French risk-free rate
url = "https://mba.tuck.dartmouth.edu/pages/faculty/ken.french/ftp/F-F_Research_Data_Factors_CSV.zip"
resp = requests.get(url, timeout=60, headers={"User-Agent": "Mozilla/5.0"})
resp.raise_for_status()
zf = zipfile.ZipFile(io.BytesIO(resp.content))
name = [n for n in zf.namelist() if n.lower().endswith(".csv")][0]
text = zf.read(name).decode("latin-1")

rf_rows = {}
for row in csv.reader(text.splitlines()):
    if not row:
        continue
    first = row[0].strip()
    if len(first) == 6 and first.isdigit():
        # columns: YYYYMM, Mkt-RF, SMB, HML, RF   (percent)
        rf_rows[first] = float(row[4]) / 100.0
rf_series = pd.Series(rf_rows)
rf_series.index = pd.to_datetime(rf_series.index, format="%Y%m") + pd.offsets.MonthEnd(0)
rf_series = rf_series.sort_index()
rf_m = rf_series.reindex(rets.index)
assert not rf_m.isna().any(), "Missing RF values for some months"
assert len(rf_m) == 48

# ----------------------------------------------------------------------------
# STATISTICS
# ----------------------------------------------------------------------------
mu = rets.mean().values * 12.0
sd_assets = rets.std(ddof=1).values * np.sqrt(12.0)
rf = float(rf_m.mean() * 12.0)
cov = rets.cov(ddof=1).values * 12.0

print("Annualised risk-free rate: %.4f%%" % (rf * 100))
print(pd.DataFrame({"ret_%": mu * 100, "sd_%": sd_assets * 100}, index=TICKERS).round(2))

# ----------------------------------------------------------------------------
# OPTIMISATION HELPERS
# ----------------------------------------------------------------------------
rng = np.random.default_rng(42)
BOUNDS = [(0.0, 1.0)] * N


def starts(n_random=14, extra=None):
    pts = [np.ones(N) / N]
    pts += [w for w in rng.dirichlet(np.ones(N), size=n_random)]
    if extra is not None:
        pts.append(extra)
    return pts  # always >= 10 starting points


def multistart(fun, cons, x0s, tol_viol=1e-6):
    best_w, best_f = None, np.inf
    fb_w, fb_v = None, np.inf  # fallback: least constraint violation
    for x0 in x0s:
        try:
            res = minimize(fun, x0, method="SLSQP", bounds=BOUNDS,
                           constraints=cons, options={"maxiter": 1000, "ftol": 1e-12})
        except Exception:
            continue
        w = np.clip(res.x, 0.0, 1.0)
        viol = 0.0
        for c in cons:
            viol = max(viol, abs(float(c["fun"](w))))
        if viol < tol_viol and np.isfinite(res.fun) and res.fun < best_f:
            best_f, best_w = res.fun, w
        if viol < fb_v:
            fb_v, fb_w = viol, w
    if best_w is None:
        best_w = fb_w
    return best_w


def port_stats(w):
    r = float(w @ mu)
    v = float(w @ cov @ w)
    return r, np.sqrt(v), v


sum_con = {"type": "eq", "fun": lambda w: np.sum(w) - 1.0}

# ----------------------------------------------------------------------------
# PORTFOLIOS
# ----------------------------------------------------------------------------
# MVP
w_mvp = multistart(lambda w: w @ cov @ w, [sum_con], starts())
w_mvp = w_mvp / w_mvp.sum()
r_mvp, s_mvp, v_mvp = port_stats(w_mvp)

# ORP (max Sharpe)
def neg_sharpe(w):
    v = w @ cov @ w
    return -((w @ mu - rf) / np.sqrt(max(v, 1e-16)))

w_orp = multistart(neg_sharpe, [sum_con], starts(n_random=20))
w_orp = w_orp / w_orp.sum()
if np.sum(w_orp > 0.0001) < 2:
    raise ValueError("ORP holds fewer than 2 ETFs with weight > 0.0001. Please choose other ETFs.")
r_orp, s_orp, v_orp = port_stats(w_orp)
sharpe_orp = (r_orp - rf) / s_orp

# OCP (A = 7)
A = 7.0
y = (r_orp - rf) / (A * v_orp)
r_ocp = rf + y * (r_orp - rf)
s_ocp = y * s_orp
w_ocp = y * w_orp
w_ocp_rf = 1.0 - y
sharpe_ocp = (r_ocp - rf) / s_ocp
sharpe_mvp = (r_mvp - rf) / s_mvp

print("\nMVP: ret %.2f%% sd %.2f%%" % (r_mvp * 100, s_mvp * 100))
print("ORP: ret %.2f%% sd %.2f%% Sharpe %.3f" % (r_orp * 100, s_orp * 100, sharpe_orp))
print("OCP: y = %.3f, ret %.2f%% sd %.2f%%" % (y, r_ocp * 100, s_ocp * 100))

# ----------------------------------------------------------------------------
# MINIMUM VARIANCE FRONTIER
# ----------------------------------------------------------------------------
targets = np.linspace(mu.min(), mu.max(), 80)
frontier = []
prev = None
for t in targets:
    cons = [sum_con, {"type": "eq", "fun": (lambda w, t=t: w @ mu - t)}]
    w = multistart(lambda w: w @ cov @ w, cons, starts(n_random=14, extra=prev))
    prev = w
    r, s, _ = port_stats(w)
    frontier.append({"target": t, "ret": r, "sd": s, "w": w})
front_df = pd.DataFrame({
    "target": [f["target"] for f in frontier],
    "ret": [f["ret"] for f in frontier],
    "sd": [f["sd"] for f in frontier],
})
W_front = np.array([f["w"] for f in frontier])

# ----------------------------------------------------------------------------
# CHART
# ----------------------------------------------------------------------------
fig, ax = plt.subplots(figsize=(11, 7.5))
ax.plot(front_df["sd"] * 100, front_df["ret"] * 100, color="tab:blue", lw=2,
        label="Minimum variance frontier (no short selling)")

ax.scatter(sd_assets * 100, mu * 100, s=25, color="grey", zorder=3, label="Individual ETFs")
for i, tk in enumerate(TICKERS):
    ax.annotate(tk, (sd_assets[i] * 100, mu[i] * 100), xytext=(4, 4),
                textcoords="offset points", fontsize=8, color="dimgrey")

# Capital Allocation Line
x_max = max(1.7 * s_orp, 1.05 * s_ocp)
xs = np.array([0.0, x_max])
ys = rf + sharpe_orp * xs
ax.plot(xs * 100, ys * 100, ls="--", color="tab:green", lw=1.6,
        label="Capital Allocation Line")

ax.scatter([s_mvp * 100], [r_mvp * 100], s=220, marker="D", color="tab:red",
           edgecolor="black", zorder=5, label="MVP")
ax.scatter([s_orp * 100], [r_orp * 100], s=320, marker="*", color="gold",
           edgecolor="black", zorder=5, label="ORP")
ax.scatter([s_ocp * 100], [r_ocp * 100], s=200, marker="s", color="tab:purple",
           edgecolor="black", zorder=5, label="OCP (A = 7)")

ax.annotate("MVP (%.2f%%, %.2f%%)" % (s_mvp * 100, r_mvp * 100),
            (s_mvp * 100, r_mvp * 100), xytext=(12, -18), textcoords="offset points",
            fontsize=9, fontweight="bold", color="tab:red",
            arrowprops=dict(arrowstyle="-", color="tab:red"))
ax.annotate("ORP (%.2f%%, %.2f%%)" % (s_orp * 100, r_orp * 100),
            (s_orp * 100, r_orp * 100), xytext=(-110, 22), textcoords="offset points",
            fontsize=9, fontweight="bold", color="darkgoldenrod",
            arrowprops=dict(arrowstyle="-", color="darkgoldenrod"))
ax.annotate("OCP (%.2f%%, %.2f%%)" % (s_ocp * 100, r_ocp * 100),
            (s_ocp * 100, r_ocp * 100), xytext=(12, -22), textcoords="offset points",
            fontsize=9, fontweight="bold", color="tab:purple",
            arrowprops=dict(arrowstyle="-", color="tab:purple"))

ax.set_xlim(left=0)
ax.set_xlabel("Annualised standard deviation (%)")
ax.set_ylabel("Annualised return (%)")
ax.set_title("Minimum Variance Frontier, Optimal Risky and Complete Portfolios\n"
             "10 ETFs, monthly data Jan 2021 - Dec 2024, no short selling")
ax.grid(True, alpha=0.3, linestyle=":")
ax.legend(loc="lower right", fontsize=9)
fig.tight_layout()
fig.savefig("mvf_ai.png", dpi=200)
plt.show()

# ----------------------------------------------------------------------------
# EXCEL
# ----------------------------------------------------------------------------
wb = Workbook()
ws = wb.active
ws.title = "MVF"
bold = Font(bold=True)
hdr_fill = PatternFill("solid", fgColor="DDEBF7")

ws["A1"] = "(a) Minimum variance frontier: 80 points (annualised, decimals)"
ws["A1"].font = bold
hdr = ["Target return", "SD"] + TICKERS
for j, h in enumerate(hdr, start=1):
    c = ws.cell(row=2, column=j, value=h)
    c.font = bold
    c.fill = hdr_fill
    c.alignment = Alignment(horizontal="center")
for i, f in enumerate(frontier):
    row = 3 + i
    ws.cell(row=row, column=1, value=float(f["target"]))
    ws.cell(row=row, column=2, value=float(f["sd"]))
    for k in range(N):
        ws.cell(row=row, column=3 + k, value=float(f["w"][k]))
    for col in range(1, 13):
        ws.cell(row=row, column=col).number_format = "0.0000"

start = 3 + 80 + 3  # leave a gap below the frontier table
ws.cell(row=start, column=1, value="(b) MVP, ORP and OCP (annualised, decimals; RF = weight in risk-free asset)").font = bold
hdr2 = ["Portfolio"] + TICKERS + ["RF", "Return", "SD", "Sharpe"]
for j, h in enumerate(hdr2, start=1):
    c = ws.cell(row=start + 1, column=j, value=h)
    c.font = bold
    c.fill = hdr_fill
    c.alignment = Alignment(horizontal="center")
rows = [
    ("MVP", w_mvp, 0.0, r_mvp, s_mvp, sharpe_mvp),
    ("ORP", w_orp, 0.0, r_orp, s_orp, sharpe_orp),
    ("OCP", w_ocp, w_ocp_rf, r_ocp, s_ocp, sharpe_ocp),
]
for i, (nm, w, wrf, r, s, sh) in enumerate(rows):
    row = start + 2 + i
    ws.cell(row=row, column=1, value=nm).font = bold
    for k in range(N):
        ws.cell(row=row, column=2 + k, value=float(w[k]))
    ws.cell(row=row, column=2 + N, value=float(wrf))
    ws.cell(row=row, column=3 + N, value=float(r))
    ws.cell(row=row, column=4 + N, value=float(s))
    ws.cell(row=row, column=5 + N, value=float(sh))
    for col in range(2, 6 + N):
        ws.cell(row=row, column=col).number_format = "0.0000"
ws.cell(row=start + 6, column=1, value="Risk-free rate (annualised)")
ws.cell(row=start + 6, column=2, value=float(rf)).number_format = "0.0000"
ws.cell(row=start + 7, column=1, value="OCP y (risky share, A = 7)")
ws.cell(row=start + 7, column=2, value=float(y)).number_format = "0.0000"

for col in range(1, 17):
    letter = ws.cell(row=2, column=col).column_letter
    ws.column_dimensions[letter].width = 13
ws.column_dimensions["A"].width = 26

img = XLImage("mvf_ai.png")
img.width, img.height = 880, 600
ws.add_image(img, "N2")
wb.save("ai_mvf.xlsx")

# ----------------------------------------------------------------------------
# JSON
# ----------------------------------------------------------------------------
def wdict(w):
    return {t: float(x) for t, x in zip(TICKERS, w)}

ocp_w = wdict(w_ocp)
ocp_w["RF"] = float(w_ocp_rf)
results = {
    "mvp": {"weights": wdict(w_mvp), "ret": float(r_mvp), "sd": float(s_mvp)},
    "orp": {"weights": wdict(w_orp), "ret": float(r_orp), "sd": float(s_orp)},
    "ocp": {"weights": ocp_w, "ret": float(r_ocp), "sd": float(s_ocp)},
}
with open("results.json", "w") as fh:
    json.dump(results, fh, indent=2)

print("\nDONE")